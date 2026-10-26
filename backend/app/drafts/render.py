"""Output formats for a Draft. Each is a pure function of the block list."""
from __future__ import annotations

import io
import textwrap

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .engine import Draft

WIDTH = 90


def to_text(draft: Draft) -> str:
    """Plain text with fixed-width alignment: good for copy-paste into an
    email or a word processor that will reformat it anyway."""
    lines: list[str] = []
    for b in draft.blocks:
        if b.kind == "spacer":
            lines.append("")
        elif b.kind in ("title", "center"):
            lines.extend(l.center(WIDTH).rstrip() for l in textwrap.wrap(b.text, WIDTH - 10) or [""])
        elif b.kind == "right":
            lines.extend(l.rjust(WIDTH) for l in textwrap.wrap(b.text, WIDTH // 2) or [""])
        elif b.kind in ("numbered", "prayer"):
            prefix = f"{b.number} "
            lines.extend(textwrap.wrap(b.text, WIDTH, initial_indent=prefix,
                                       subsequent_indent=" " * len(prefix)))
            lines.append("")
        else:
            lines.extend(textwrap.wrap(b.text, WIDTH) or [""])
            if b.kind == "para":
                lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def to_docx(draft: Draft) -> bytes:
    """Court-style Word document.

    Formatting follows common Indian court filing practice: A4, Times New
    Roman 12pt, 1.5 line spacing, justified body text, and a wider left margin
    (1.5") to leave room for binding/stitching on the left edge.
    """
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)  # A4
    sec.left_margin, sec.right_margin = Inches(1.5), Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    align = {
        "title": WD_ALIGN_PARAGRAPH.CENTER,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "para": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "numbered": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "prayer": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }
    for b in draft.blocks:
        if b.kind == "spacer":
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        p.alignment = align.get(b.kind, WD_ALIGN_PARAGRAPH.LEFT)
        if b.kind in ("numbered", "prayer"):
            # Hanging indent so wrapped lines align with the text, not the number.
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.add_run(f"{b.number}\t")
        run = p.add_run(b.text)
        run.bold = b.bold or b.kind == "title"
        if b.kind == "title":
            run.font.size = Pt(13)
            run.underline = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
